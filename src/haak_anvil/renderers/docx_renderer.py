"""DOCX renderer — professional Word pentest report via python-docx.

Produces a client-ready ``.docx``: cover, executive summary with a severity
table, and one section per finding (critical → info). Requires the optional
``docx`` extra::

    pip install "haak-anvil[docx]"

DOCX is a binary format, so :meth:`render` (which must return ``str``) raises;
callers use :meth:`write` to emit the file.
"""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

from haak_anvil.core.models import Finding, ReportBundle
from haak_anvil.core.severity import Severity
from haak_anvil.renderers.base import RendererBase

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument

_SEV_RGB = {
    Severity.CRITICAL: (0xC0, 0x1B, 0x1B),
    Severity.HIGH: (0xE8, 0x6A, 0x17),
    Severity.MEDIUM: (0xC9, 0xA2, 0x00),
    Severity.LOW: (0x1F, 0x6F, 0xB2),
    Severity.INFO: (0x6B, 0x72, 0x80),
}


class DocxRenderer(RendererBase):
    """Render a ReportBundle into a Microsoft Word document."""

    extension = "docx"

    def render(self, bundle: ReportBundle) -> str:
        raise NotImplementedError(
            "DOCX is binary; use DocxRenderer().write(bundle, path) instead of render()."
        )

    def write(self, bundle: ReportBundle, path: Path | str) -> Path:
        path = Path(path)
        if path.is_dir() or path.suffix == "":
            path = path / f"{bundle.engagement.id}.{self.extension}"
        path.parent.mkdir(parents=True, exist_ok=True)
        doc = self._build(bundle)
        doc.save(str(path))
        return path

    # ------------------------------------------------------------------ builder

    def _build(self, bundle: ReportBundle) -> DocxDocument:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Pt, RGBColor
        except ModuleNotFoundError as exc:  # pragma: no cover - import guard
            raise RuntimeError(
                "python-docx is required for DOCX output. Install with: "
                'pip install "haak-anvil[docx]"'
            ) from exc

        eng = bundle.engagement
        doc = Document()

        # ----- cover -----
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("Reporte de Pruebas de Penetración")
        run.bold = True
        run.font.size = Pt(24)

        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = subtitle.add_run(eng.client_name)
        sub_run.font.size = Pt(16)
        sub_run.font.color.rgb = RGBColor(0x00, 0xA8, 0xCC)

        meta = doc.add_table(rows=0, cols=2)
        meta.style = "Light List Accent 1"
        rows = [
            ("Engagement", eng.id),
            ("Alcance", eng.scope),
            ("Metodología", eng.methodology),
            ("Analista", eng.analyst),
            ("Periodo", self._period(eng)),
            ("Generado", bundle.generated_at.strftime("%Y-%m-%d %H:%M UTC")),
            ("Herramienta", f"{bundle.generator} {bundle.generator_version}"),
        ]
        for label, value in rows:
            cells = meta.add_row().cells
            cells[0].paragraphs[0].add_run(label).bold = True
            cells[1].text = str(value)

        # ----- executive summary -----
        doc.add_page_break()
        doc.add_heading("Resumen Ejecutivo", level=1)
        doc.add_paragraph(
            f"Se identificaron {len(bundle.findings)} hallazgos sobre "
            f"{len(bundle.assets)} activos evaluados. La distribución por severidad "
            f"se muestra a continuación."
        )

        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].paragraphs[0].add_run("Severidad").bold = True
        hdr[1].paragraphs[0].add_run("Hallazgos").bold = True
        for sev in sorted(Severity, key=lambda s: -s.numeric):
            count = bundle.severity_breakdown[sev.value]
            cells = table.add_row().cells
            run = cells[0].paragraphs[0].add_run(sev.label_es)
            run.bold = True
            run.font.color.rgb = RGBColor(*_SEV_RGB[sev])
            cells[1].text = str(count)

        # ----- findings detail -----
        doc.add_page_break()
        doc.add_heading("Hallazgos", level=1)
        ordered = bundle.findings_sorted()
        if not ordered:
            doc.add_paragraph("Sin hallazgos registrados.")
        for i, finding in enumerate(ordered, start=1):
            self._add_finding(doc, i, finding)

        return doc

    # ------------------------------------------------------------------ helpers

    @staticmethod
    def _period(eng) -> str:
        if eng.period_start and eng.period_end:
            return f"{eng.period_start} — {eng.period_end}"
        if eng.period_start:
            return str(eng.period_start)
        return "N/A"

    def _add_finding(self, doc, index: int, f: Finding) -> None:
        from docx.shared import RGBColor

        heading = doc.add_heading(level=2)
        run = heading.add_run(f"{index}. {f.title}")
        run.font.color.rgb = RGBColor(*_SEV_RGB[f.severity])

        badge = doc.add_paragraph()
        b = badge.add_run(f"Severidad: {f.severity.label_es}")
        b.bold = True
        b.font.color.rgb = RGBColor(*_SEV_RGB[f.severity])
        if f.cvss is not None:
            badge.add_run(f"    ·    CVSS {f.cvss.version}: {f.cvss.score}")
        if f.epss_score is not None:
            pct = f" (pctl {f.epss_percentile:.2%})" if f.epss_percentile is not None else ""
            badge.add_run(f"    ·    EPSS: {f.epss_score:.2%}{pct}")

        detail = doc.add_table(rows=0, cols=2)
        detail.style = "Light List"
        meta_rows = [
            ("Activo", f.asset or "—"),
            ("Puerto", f"{f.port}/{f.protocol}" if f.port else "—"),
            ("CVE", ", ".join(f.cve) or "—"),
            ("CWE", ", ".join(f.cwe) or "—"),
            ("Fuente", f"{f.tool}" + (f" ({f.plugin_id})" if f.plugin_id else "")),
        ]
        for label, value in meta_rows:
            cells = detail.add_row().cells
            cells[0].paragraphs[0].add_run(label).bold = True
            cells[1].text = str(value)

        if f.description:
            doc.add_heading("Descripción", level=3)
            doc.add_paragraph(f.description)
        if f.impact:
            doc.add_heading("Impacto", level=3)
            doc.add_paragraph(f.impact)
        if f.evidence:
            doc.add_heading("Evidencia", level=3)
            ev = doc.add_paragraph(f.evidence)
            ev.style = doc.styles["Intense Quote"] if "Intense Quote" in [s.name for s in doc.styles] else ev.style
        if f.remediation:
            doc.add_heading("Remediación", level=3)
            doc.add_paragraph(f.remediation)
        if f.references:
            doc.add_heading("Referencias", level=3)
            for ref in f.references:
                doc.add_paragraph(ref, style="List Bullet")
