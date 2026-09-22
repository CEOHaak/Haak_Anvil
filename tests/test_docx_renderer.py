import pytest

from haak_anvil.core.models import Finding, ReportBundle
from haak_anvil.core.severity import CVSS, Severity
from haak_anvil.renderers import DocxRenderer, get_renderer

pytest.importorskip("docx", reason="python-docx not installed (optional 'docx' extra)")


def _bundle(engagement) -> ReportBundle:
    return ReportBundle(
        engagement=engagement,
        findings=[
            Finding(
                id="f1",
                title="Log4Shell RCE",
                severity=Severity.CRITICAL,
                cvss=CVSS(score=10.0, vector="CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:C/C:H/I:H/A:H"),
                description="Remote code execution via JNDI lookup.",
                remediation="Upgrade Log4j.",
                references=["https://logging.apache.org"],
                cve=["CVE-2021-44228"],
                cwe=["CWE-502"],
                asset="acme.example",
                port=8443,
                protocol="tcp",
                epss_score=0.975,
                epss_percentile=0.999,
                tool="nuclei",
            ),
            Finding(id="f2", title="Info banner", severity=Severity.INFO, tool="burp"),
        ],
    )


def test_docx_factory_registered():
    assert get_renderer("docx") is DocxRenderer
    assert get_renderer("word") is DocxRenderer


def test_docx_render_raises():
    with pytest.raises(NotImplementedError):
        DocxRenderer().render(_bundle_engagement())


def _bundle_engagement():
    from haak_anvil.core.engagement import Engagement

    return _bundle(Engagement(id="T-1", client_name="Acme", scope="s"))


def test_docx_write_produces_readable_file(tmp_path, engagement):
    from docx import Document

    out = tmp_path / "report.docx"
    written = DocxRenderer().write(_bundle(engagement), out)
    assert written == out
    assert out.exists() and out.stat().st_size > 0

    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert engagement.client_name in text
    assert "Log4Shell RCE" in text
    # Table cells hold severity + metadata; make sure the doc has tables
    assert len(doc.tables) >= 2


def test_docx_write_infers_name_from_dir(tmp_path, engagement):
    written = DocxRenderer().write(_bundle(engagement), tmp_path)
    assert written.name == f"{engagement.id}.docx"
    assert written.exists()
